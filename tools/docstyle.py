"""docstyle.py - utilidades de diseño para los Word del Lab 3 (python-docx)."""
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

NAVY = RGBColor(0x0B, 0x2A, 0x4A)
BLUE = RGBColor(0x1F, 0x5F, 0xA8)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
GREY = RGBColor(0x5B, 0x65, 0x70)
DARK = RGBColor(0x1E, 0x29, 0x33)
FONT = "Calibri"


def _shade(el, hex_fill):
    pr = el.get_or_add_tcPr() if hasattr(el, "get_or_add_tcPr") else el
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def _cell_borders(cell, color="D0D7DE", sz="4", sides=("top", "left", "bottom", "right")):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for s in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{s}")
        if s in sides:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), sz); e.set(qn("w:color"), color)
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tcPr.append(b)


def _margins(cell, top=90, bottom=90, left=130, right=130):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)


def _keep_together(table):
    """Evita que filas/tablas se partan entre páginas."""
    rows = table.rows
    for i, row in enumerate(rows):
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
        if i < len(rows) - 1:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True


def _field(run, instr):
    for t, v in (("begin", None), (None, instr), ("end", None)):
        if t:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), t); run._r.append(f)
        else:
            i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = v; run._r.append(i)


class Doc:
    def __init__(self, title_short):
        self.d = Document()
        self.fig = 0
        self.tab = 0
        s = self.d.sections[0]
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.top_margin, s.bottom_margin = Mm(20), Mm(18)
        s.left_margin = s.right_margin = Mm(20)
        st = self.d.styles
        n = st["Normal"]; n.font.name = FONT; n.font.size = Pt(10.5); n.font.color.rgb = DARK
        n.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        n.paragraph_format.space_after = Pt(6); n.paragraph_format.line_spacing = 1.12
        for lvl, size, color, before in ((1, 17, NAVY, 16), (2, 13, BLUE, 12), (3, 11.5, TEAL, 8)):
            h = st[f"Heading {lvl}"]
            h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = color
            h.element.rPr.rFonts.set(qn("w:asciiTheme"), FONT) if False else None
            rf = h.element.rPr.rFonts
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), FONT)
            h.paragraph_format.space_before = Pt(before); h.paragraph_format.space_after = Pt(5)
            h.paragraph_format.keep_with_next = True
        self.title_short = title_short

    # ---------- portada ----------
    def cover(self, kicker, title, subtitle, meta_rows, note):
        d = self.d
        band = d.add_table(rows=1, cols=1); band.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = band.rows[0].cells[0]; _shade(c._tc.get_or_add_tcPr(), "0B2A4A"); _cell_borders(c, sides=())
        _margins(c, 500, 500, 420, 420)
        p = c.paragraphs[0]; r = p.add_run(kicker.upper()); r.font.size = Pt(10); r.bold = True
        r.font.color.rgb = RGBColor(0x9C, 0xC3, 0xF0)
        p = c.add_paragraph(); r = p.add_run(title); r.font.size = Pt(28); r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); p.paragraph_format.space_before = Pt(8)
        p = c.add_paragraph(); r = p.add_run(subtitle); r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0xD6, 0xE4, 0xF5)
        d.add_paragraph().paragraph_format.space_after = Pt(18)
        t = d.add_table(rows=0, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, v in meta_rows:
            row = t.add_row().cells
            for cell in row:
                _cell_borders(cell, color="D0D7DE", sides=("bottom",)); _margins(cell, 110, 110, 120, 120)
            row[0].width, row[1].width = Mm(55), Mm(115)
            rr = row[0].paragraphs[0].add_run(k); rr.bold = True; rr.font.color.rgb = GREY; rr.font.size = Pt(10)
            rv = row[1].paragraphs[0].add_run(v); rv.font.size = Pt(10.5)
        d.add_paragraph().paragraph_format.space_after = Pt(24)
        self.callout(note, kind="info")
        d.add_page_break()
        self._header_footer()

    def _header_footer(self):
        s = self.d.sections[0]
        s.different_first_page_header_footer = True
        hp = s.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run(self.title_short); r.font.size = Pt(8.5); r.font.color.rgb = GREY
        fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for txt, instr in (("Página ", "PAGE"), (" de ", "NUMPAGES")):
            r = fp.add_run(txt); r.font.size = Pt(8.5); r.font.color.rgb = GREY
            r = fp.add_run(); r.font.size = Pt(8.5); r.font.color.rgb = GREY; _field(r, instr)

    def toc(self):
        tp = self.d.add_paragraph(); r = tp.add_run("Contenido"); r.bold = True
        r.font.size = Pt(17); r.font.color.rgb = NAVY; tp.paragraph_format.space_after = Pt(10)
        p = self.d.add_paragraph(); _field(p.add_run(), 'TOC \\o "1-2" \\h \\z \\u')
        self.d.add_page_break()

    # ---------- bloques ----------
    def h(self, text, lvl=1):
        return self.d.add_heading(text, level=lvl)

    def p(self, *parts, size=None, align=None, after=None):
        """parts: str o (texto, {'b':1,'i':1,'code':1,'color':RGB})"""
        para = self.d.add_paragraph()
        for part in parts:
            txt, fmt = (part, {}) if isinstance(part, str) else part
            r = para.add_run(txt)
            r.bold = bool(fmt.get("b")); r.italic = bool(fmt.get("i"))
            if fmt.get("code"):
                r.font.name = "Consolas"; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0x9A, 0x34, 0x12)
            if fmt.get("color"):
                r.font.color.rgb = fmt["color"]
            if size:
                r.font.size = Pt(size)
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        return para

    def bullets(self, items):
        for it in items:
            para = self.d.add_paragraph(style="List Bullet")
            parts = it if isinstance(it, (list, tuple)) else [it]
            for part in parts:
                txt, fmt = (part, {}) if isinstance(part, str) else part
                r = para.add_run(txt); r.bold = bool(fmt.get("b"))
                if fmt.get("code"):
                    r.font.name = "Consolas"; r.font.size = Pt(9.5)
            para.paragraph_format.space_after = Pt(2)

    def callout(self, text, kind="info", title=None):
        palette = {"info": ("EAF2FC", "1F5FA8"), "ok": ("E8F6EF", "1E8449"),
                   "warn": ("FFF6E5", "C47A19"), "key": ("F1ECFA", "6B3FA0")}
        fill, bar = palette[kind]
        t = self.d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]; _shade(c._tc.get_or_add_tcPr(), fill)
        _cell_borders(c, color=bar, sz="24", sides=("left",)); _margins(c, 120, 120, 200, 160)
        para = c.paragraphs[0]
        if title:
            r = para.add_run(title + "  "); r.bold = True; r.font.color.rgb = RGBColor.from_string(bar)
        parts = text if isinstance(text, (list, tuple)) else [text]
        for part in parts:
            txt, fmt = (part, {}) if isinstance(part, str) else part
            r = para.add_run(txt); r.bold = bool(fmt.get("b")); r.font.size = Pt(10)
            if fmt.get("code"):
                r.font.name = "Consolas"; r.font.size = Pt(9.5)
        _keep_together(t)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def code(self, lines):
        t = self.d.add_table(rows=1, cols=1)
        c = t.rows[0].cells[0]; _shade(c._tc.get_or_add_tcPr(), "0F172A")
        _cell_borders(c, sides=()); _margins(c, 120, 120, 180, 180)
        first = True
        for ln in lines:
            para = c.paragraphs[0] if first else c.add_paragraph(); first = False
            para.paragraph_format.space_after = Pt(0); para.paragraph_format.line_spacing = 1.0
            r = para.add_run(ln or " "); r.font.name = "Consolas"; r.font.size = Pt(8.8)
            r.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
        _keep_together(t)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def caption(self, prefix, text):
        para = self.d.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = para.add_run(prefix); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = BLUE
        r = para.add_run(text); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = GREY
        para.paragraph_format.space_after = Pt(10)

    def figure(self, path, caption, width_mm=165):
        para = self.d.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True; para.paragraph_format.space_after = Pt(3)
        para.add_run().add_picture(path, width=Mm(width_mm))
        self.fig += 1
        self.caption(f"Figura {self.fig}. ", caption)

    def table(self, header, rows, widths_mm=None, caption=None, align_center_cols=(), highlight=None):
        if caption:
            self.tab += 1
            cp = self.d.add_paragraph(); cp.paragraph_format.keep_with_next = True
            cp.paragraph_format.space_after = Pt(3)
            r = cp.add_run(f"Tabla {self.tab}. "); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = BLUE
            r = cp.add_run(caption); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = GREY
        t = self.d.add_table(rows=1, cols=len(header)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = t.rows[0].cells
        for i, htxt in enumerate(header):
            _shade(hdr[i]._tc.get_or_add_tcPr(), "0B2A4A"); _cell_borders(hdr[i], color="0B2A4A")
            _margins(hdr[i]); hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            r = hdr[i].paragraphs[0].add_run(htxt); r.bold = True; r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for ci, val in enumerate(row):
                c = cells[ci]; _cell_borders(c); _margins(c)
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                fill = "F4F7FB" if ri % 2 else "FFFFFF"
                if highlight and highlight(ri, ci, val):
                    fill = highlight(ri, ci, val)
                _shade(c._tc.get_or_add_tcPr(), fill)
                para = c.paragraphs[0]; para.paragraph_format.space_after = Pt(0)
                if ci in align_center_cols:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                txt, fmt = (val, {}) if isinstance(val, str) else val
                r = para.add_run(txt); r.font.size = Pt(9.5)
                r.bold = bool(fmt.get("b")) or ci == 0
                if fmt.get("code"):
                    r.font.name = "Consolas"; r.font.size = Pt(8.8)
        if widths_mm:
            for row in t.rows:
                for i, w in enumerate(widths_mm):
                    row.cells[i].width = Mm(w)
        _keep_together(t)
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    def kpis(self, items):
        """items: [(valor, etiqueta), ...] -> fila de tarjetas."""
        t = self.d.add_table(rows=1, cols=len(items)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (val, lab) in enumerate(items):
            c = t.rows[0].cells[i]; _shade(c._tc.get_or_add_tcPr(), "EAF2FC")
            _cell_borders(c, color="FFFFFF", sz="18"); _margins(c, 140, 140, 100, 100)
            p1 = c.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p1.add_run(val); r.bold = True; r.font.size = Pt(17); r.font.color.rgb = NAVY
            p1.paragraph_format.space_after = Pt(0)
            p2 = c.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p2.add_run(lab); r.font.size = Pt(8.5); r.font.color.rgb = GREY
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)

    def page_break(self):
        self.d.add_page_break()

    def save(self, path):
        self.d.save(path)
